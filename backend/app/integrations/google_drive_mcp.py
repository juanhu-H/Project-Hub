from __future__ import annotations

import io
import re
from typing import Any

import httpx
from docx import Document
from openpyxl import load_workbook

from .google_oauth import get_access_token

DRIVE_API_BASE = "https://www.googleapis.com/drive/v3"

_SEARCH_FIELDS = "nextPageToken,files(id,name,mimeType,webViewLink,parents,modifiedTime)"

_EXPORT_MIME_TYPES = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.spreadsheet": "text/csv",
    "application/vnd.google-apps.presentation": "text/plain",
}

_DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

XLSX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_XLSX_MIME_TYPE = XLSX_MIME_TYPE


def _extract_docx_text(raw: bytes) -> str:
    document = Document(io.BytesIO(raw))
    return "\n".join(
        paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
    )


def _extract_xlsx_text(raw: bytes) -> str:
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_xlsx_rows(raw: bytes) -> list[dict[str, Any]]:
    """Lee un XLSX como filas estructuradas (encabezado + datos), para no
    mezclar todos los registros de la hoja en un único bloque de texto."""
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        headers = [
            str(cell).strip() if cell is not None else f"col{i}"
            for i, cell in enumerate(header)
        ]
        for raw_row in rows_iter:
            if all(cell is None for cell in raw_row):
                continue
            record = {
                headers[i]: raw_row[i]
                for i in range(min(len(headers), len(raw_row)))
                if raw_row[i] is not None
            }
            if record:
                record["_sheet"] = sheet.title
                rows.append(record)
    return rows


def _translate_query(query: str) -> str:
    """Traduce la mini sintaxis usada en este proyecto (compatible con la que
    exponía el MCP de Drive) a la sintaxis real de `q` de Drive API v3."""

    translated = query
    translated = re.sub(r"parentId\s*!=\s*'([^']*)'", r"not '\1' in parents", translated)
    translated = re.sub(r"parentId\s*=\s*'([^']*)'", r"'\1' in parents", translated)
    translated = re.sub(r"owner\s*=\s*'([^']*)'", r"'\1' in owners", translated)
    translated = re.sub(r"\btitle\b", "name", translated)
    return translated


class GoogleDriveMCPClient:
    """Cliente que habla directo con Google Drive API v3 (REST).

    No usa el servidor MCP hosteado de Google (drivemcp.googleapis.com):
    ese servicio está en Developer Preview y, para cuentas de Workspace,
    requiere que un administrador habilite Drive para la organización y
    confíe explícitamente en la app OAuth para scopes restringidos. Llamar
    la Drive API v3 directamente evita esa dependencia y funciona igual
    con cuentas personales o de Workspace.
    """

    async def list_tools(self) -> list[dict[str, Any]]:
        return [
            {
                "name": "search_files",
                "description": "Busca archivos en Google Drive (Drive API v3 files.list).",
                "input_schema": None,
            },
            {
                "name": "read_file_content",
                "description": "Lee el contenido de texto de un archivo de Drive.",
                "input_schema": None,
            },
            {
                "name": "read_file_rows",
                "description": "Lee un archivo XLSX de Drive como filas estructuradas (encabezado + datos).",
                "input_schema": None,
            },
        ]

    async def call_tool(
        self, tool_name: str, arguments: dict[str, Any]
    ) -> dict[str, Any]:
        if tool_name == "search_files":
            return await self._search_files(arguments)
        if tool_name == "read_file_content":
            return await self._read_file_content(arguments)
        if tool_name == "read_file_rows":
            return await self._read_file_rows(arguments)
        raise ValueError(f"La herramienta '{tool_name}' no está implementada.")

    async def _auth_headers(self) -> dict[str, str]:
        token = await get_access_token()
        return {"Authorization": f"Bearer {token}"}

    async def _search_files(self, arguments: dict[str, Any]) -> dict[str, Any]:
        params: dict[str, Any] = {
            "q": _translate_query(arguments.get("query", "")),
            "pageSize": arguments.get("pageSize", 100),
            "fields": _SEARCH_FIELDS,
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
        if arguments.get("pageToken"):
            params["pageToken"] = arguments["pageToken"]

        async with httpx.AsyncClient(timeout=30) as client:
            response = await client.get(
                f"{DRIVE_API_BASE}/files",
                params=params,
                headers=await self._auth_headers(),
            )

        if response.status_code >= 400:
            raise RuntimeError(
                f"Drive API error {response.status_code}: {response.text}"
            )

        return {"tool": "search_files", "text": [], "structured": response.json()}

    async def _read_file_content(self, arguments: dict[str, Any]) -> dict[str, Any]:
        file_id = arguments["fileId"]

        async with httpx.AsyncClient(timeout=30) as client:
            headers = await self._auth_headers()

            meta_response = await client.get(
                f"{DRIVE_API_BASE}/files/{file_id}",
                params={"fields": "id,name,mimeType", "supportsAllDrives": "true"},
                headers=headers,
            )
            if meta_response.status_code >= 400:
                raise RuntimeError(
                    f"Drive API error {meta_response.status_code}: {meta_response.text}"
                )

            mime_type = meta_response.json().get("mimeType", "")
            export_mime = _EXPORT_MIME_TYPES.get(mime_type)

            if export_mime:
                content_response = await client.get(
                    f"{DRIVE_API_BASE}/files/{file_id}/export",
                    params={"mimeType": export_mime},
                    headers=headers,
                )
                if content_response.status_code >= 400:
                    raise RuntimeError(
                        f"Drive API error {content_response.status_code}: {content_response.text}"
                    )
                text = content_response.text
            elif mime_type == _DOCX_MIME_TYPE:
                content_response = await client.get(
                    f"{DRIVE_API_BASE}/files/{file_id}",
                    params={"alt": "media", "supportsAllDrives": "true"},
                    headers=headers,
                )
                if content_response.status_code >= 400:
                    raise RuntimeError(
                        f"Drive API error {content_response.status_code}: {content_response.text}"
                    )
                text = _extract_docx_text(content_response.content)
            elif mime_type == _XLSX_MIME_TYPE:
                content_response = await client.get(
                    f"{DRIVE_API_BASE}/files/{file_id}",
                    params={"alt": "media", "supportsAllDrives": "true"},
                    headers=headers,
                )
                if content_response.status_code >= 400:
                    raise RuntimeError(
                        f"Drive API error {content_response.status_code}: {content_response.text}"
                    )
                text = _extract_xlsx_text(content_response.content)
            elif mime_type.startswith("text/"):
                content_response = await client.get(
                    f"{DRIVE_API_BASE}/files/{file_id}",
                    params={"alt": "media", "supportsAllDrives": "true"},
                    headers=headers,
                )
                if content_response.status_code >= 400:
                    raise RuntimeError(
                        f"Drive API error {content_response.status_code}: {content_response.text}"
                    )
                text = content_response.text
            else:
                raise RuntimeError(
                    f"Tipo de archivo no soportado para lectura de texto: {mime_type}"
                )

        return {
            "tool": "read_file_content",
            "text": [text],
            "structured": None,
        }

    async def _read_file_rows(self, arguments: dict[str, Any]) -> dict[str, Any]:
        file_id = arguments["fileId"]

        async with httpx.AsyncClient(timeout=30) as client:
            headers = await self._auth_headers()

            meta_response = await client.get(
                f"{DRIVE_API_BASE}/files/{file_id}",
                params={"fields": "id,name,mimeType", "supportsAllDrives": "true"},
                headers=headers,
            )
            if meta_response.status_code >= 400:
                raise RuntimeError(
                    f"Drive API error {meta_response.status_code}: {meta_response.text}"
                )

            mime_type = meta_response.json().get("mimeType", "")
            if mime_type != XLSX_MIME_TYPE:
                raise RuntimeError(
                    f"read_file_rows solo soporta XLSX, recibido: {mime_type}"
                )

            content_response = await client.get(
                f"{DRIVE_API_BASE}/files/{file_id}",
                params={"alt": "media", "supportsAllDrives": "true"},
                headers=headers,
            )
            if content_response.status_code >= 400:
                raise RuntimeError(
                    f"Drive API error {content_response.status_code}: {content_response.text}"
                )

        rows = _extract_xlsx_rows(content_response.content)
        return {
            "tool": "read_file_rows",
            "text": [],
            "structured": {"rows": rows},
        }
